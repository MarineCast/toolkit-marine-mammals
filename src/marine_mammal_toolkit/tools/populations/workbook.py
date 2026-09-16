"""Reusable annual census workbook validation and reconciliation."""

from __future__ import annotations
import math
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from openpyxl import load_workbook


def _normalize_header(value: object, aliases: dict[str, str]) -> str:
    """Normalize a workbook header to a stable snake-case field name."""
    text = " ".join(str(value or "").strip().lower().replace("_", " ").split())
    return aliases.get(text, text.replace(" ", "_"))


def _validate_headers(
    headers: list[str], *, sheet_name: str, required_columns: set[str]
) -> None:
    """Validate required and duplicate normalized workbook headers."""
    duplicate_headers = sorted(
        header for header, count in Counter(headers).items() if header and count > 1
    )
    if duplicate_headers:
        raise ValueError(
            f"Workbook sheet {sheet_name!r} contains duplicate normalized "
            f"columns: {duplicate_headers}."
        )

    missing = required_columns - set(headers)
    if missing:
        raise KeyError(
            f"Workbook sheet {sheet_name!r} is missing required columns: "
            f"{sorted(missing)}."
        )


def _to_int(
    value: object,
    *,
    column: str,
    row_number: int,
    minimum: int | None = None,
    maximum: int | None = None,
) -> int:
    """Convert a workbook cell to an integer without silently truncating it."""
    if value is None or str(value).strip() == "":
        raise ValueError(
            f"Missing {column!r} at workbook row {row_number}. "
            "If the cell contains a formula, recalculate and save the workbook "
            "before running the export."
        )

    try:
        number = float(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"Invalid number for {column!r} at workbook row "
            f"{row_number}: {value!r}."
        ) from exc

    if not math.isfinite(number) or not number.is_integer():
        raise ValueError(
            f"Expected a whole number for {column!r} at workbook row "
            f"{row_number}, got {value!r}."
        )

    result = int(number)
    if minimum is not None and result < minimum:
        raise ValueError(
            f"Value for {column!r} at workbook row {row_number} must be at "
            f"least {minimum}, got {result}."
        )
    if maximum is not None and result > maximum:
        raise ValueError(
            f"Value for {column!r} at workbook row {row_number} must be at "
            f"most {maximum}, got {result}."
        )
    return result


def read_annual_counts(
    source_path: str | Path,
    *,
    sheet_name: str,
    aliases: dict[str, str],
    count_columns: tuple[str, ...],
    component_columns: tuple[str, ...],
    total_column: str,
) -> list[dict[str, Any]]:
    """Load, validate, and normalize SRKW population rows from a workbook."""
    source = Path(source_path).expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(f"Population workbook does not exist: {source}")
    if not source.is_file():
        raise ValueError(f"Population source is not a file: {source}")

    workbook = load_workbook(source, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise KeyError(
                f"Workbook {source} does not contain sheet {sheet_name!r}. "
                f"Available sheets: {workbook.sheetnames}."
            )

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError(
                f"Workbook sheet {sheet_name!r} is empty: {source}"
            ) from exc

        headers = [_normalize_header(value, aliases) for value in raw_headers]
        _validate_headers(
            headers,
            sheet_name=sheet_name,
            required_columns={"census_year", *count_columns},
        )

        normalized: list[dict[str, Any]] = []
        seen_years: set[int] = set()
        max_reasonable_year = datetime.now(timezone.utc).year + 1

        for row_number, raw_row in enumerate(rows, start=2):
            raw = dict(zip(headers, raw_row))
            if all(value is None or str(value).strip() == "" for value in raw.values()):
                continue

            census_year = _to_int(
                raw.get("census_year"),
                column="census_year",
                row_number=row_number,
                minimum=1900,
                maximum=max_reasonable_year,
            )
            if census_year in seen_years:
                raise ValueError(
                    f"Duplicate census year {census_year} at workbook row "
                    f"{row_number}."
                )
            seen_years.add(census_year)

            values = {
                column: _to_int(
                    raw.get(column),
                    column=column,
                    row_number=row_number,
                    minimum=0,
                )
                for column in count_columns
            }
            pod_sum = sum(values[column] for column in component_columns)

            record: dict[str, Any] = {
                "census_year": census_year,
                **values,
                "pod_sum": pod_sum,
                "total_matches_pod_sum": pod_sum == values[total_column],
            }
            normalized.append(record)
    finally:
        workbook.close()

    normalized.sort(key=lambda item: item["census_year"])
    if not normalized:
        raise ValueError(f"No population rows found in {source} sheet {sheet_name!r}.")
    return normalized
